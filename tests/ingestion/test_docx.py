"""Tests for the DOCX loader."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as _DocxDocument

from gyroscope.core.models import DocumentKind
from gyroscope.ingestion.docx import DocxLoader


def _make_fixture(path: Path) -> None:
    d = _DocxDocument()
    d.add_heading("Top Heading", level=1)
    d.add_paragraph("Intro paragraph one.")
    d.add_heading("Section A", level=2)
    d.add_paragraph("Body of section A.")
    d.add_heading("Subsection A.1", level=3)
    d.add_paragraph("More content.")
    d.save(str(path))


@pytest.mark.asyncio
async def test_docx_loader_headings(tmp_path: Path) -> None:
    path = tmp_path / "fixture.docx"
    _make_fixture(path)

    loader = DocxLoader()
    assert loader.can_load(str(path))
    assert not loader.can_load("http://x.com/a.docx")

    docs = await loader.load(str(path))
    assert len(docs) == 1
    doc = docs[0]
    assert doc.kind == DocumentKind.DOCX
    assert "Body of section A." in doc.text
    assert "More content." in doc.text
    headings = doc.metadata["headings"]
    texts = [h["text"] for h in headings]
    levels = [h["level"] for h in headings]
    assert "Top Heading" in texts
    assert "Section A" in texts
    assert "Subsection A.1" in texts
    assert 1 in levels and 2 in levels and 3 in levels
    assert doc.title == "Top Heading"
