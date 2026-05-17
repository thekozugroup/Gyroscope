"""DOCX loader using python-docx.

Paragraphs are concatenated in document order; ``Heading N`` styles are
captured in ``metadata["headings"]`` so structure survives ingestion.
"""

from __future__ import annotations

import asyncio
import re
from pathlib import Path
from typing import Any

from docx import Document as _DocxDocument

from gyroscope.core.logging import get_logger
from gyroscope.core.models import Document, DocumentKind
from gyroscope.ingestion.base import Loader, is_url

logger = get_logger(__name__)

_HEADING_STYLE_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


def _heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    m = _HEADING_STYLE_RE.match(style_name.strip())
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            return None
    if style_name.strip().lower() == "title":
        return 0
    return None


def _parse_docx(path: Path) -> tuple[str, str | None, list[dict[str, Any]]]:
    docx = _DocxDocument(str(path))
    lines: list[str] = []
    headings: list[dict[str, Any]] = []
    title: str | None = None
    for para in docx.paragraphs:
        text = (para.text or "").strip()
        style_name = getattr(getattr(para, "style", None), "name", None)
        level = _heading_level(style_name)
        if level is not None and text:
            headings.append({"level": level, "text": text})
            if level == 0 and title is None:
                title = text
            lines.append(("#" * max(1, level) if level > 0 else "#") + " " + text)
        elif text:
            lines.append(text)
    body = "\n\n".join(lines).strip() + "\n"
    if title is None:
        for h in headings:
            if h["level"] == 1:
                title = str(h["text"])
                break
    if title is None:
        core = docx.core_properties
        meta_title = getattr(core, "title", None)
        if meta_title:
            title = meta_title.strip() or None
    return body, title, headings


class DocxLoader(Loader):
    """Loader for ``.docx`` files."""

    name = "docx"
    extensions = (".docx",)

    def can_load(self, source: str) -> bool:
        if is_url(source):
            return False
        return source.lower().endswith(self.extensions)

    async def load(self, source: str) -> list[Document]:
        path = Path(source)
        text, title, headings = await asyncio.to_thread(_parse_docx, path)
        doc = Document(
            source=str(path.resolve()),
            kind=DocumentKind.DOCX,
            text=text,
            title=title,
            metadata={
                "headings": headings,
                "size_bytes": len(text.encode("utf-8")),
            },
        )
        return [doc]
